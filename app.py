import datetime
import json
import os
import re
import sys
import time
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_from_directory
from werkzeug.utils import secure_filename

load_dotenv()

BASE_DIR = Path(__file__).parent


def resource_path(relative: str) -> Path:
    """Locates a bundled, read-only asset (templates, static files) whether running as a
    normal script or as a PyInstaller-frozen exe, which extracts bundled data to a temp
    folder at sys._MEIPASS on every launch."""
    base = Path(sys._MEIPASS) if hasattr(sys, "_MEIPASS") else BASE_DIR
    return base / relative


def get_data_dir() -> Path:
    """Locates where actual content (memory, drafts, archive, generated images, the API-key
    config) should live. When running as a frozen exe, this must NOT be inside the temp
    extraction folder — that folder is wiped and rebuilt on every launch, which would silently
    lose everything the person generated last time. Instead it uses a permanent per-user
    folder (AppData on Windows, home directory elsewhere). When running as a normal script
    (development), it stays as the project folder itself, unchanged from before."""
    if hasattr(sys, "_MEIPASS"):
        appdata = os.environ.get("APPDATA")  # Windows: C:\Users\<name>\AppData\Roaming
        base = Path(appdata) if appdata else Path.home()
        d = base / "PebbleStudio"
    else:
        d = BASE_DIR
    d.mkdir(parents=True, exist_ok=True)
    return d


DATA_DIR = get_data_dir()
MEMORY_DIR = DATA_DIR / "memory"
ALLOWED_AGENTS = {"seo", "script", "image"}
ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf"}
MAX_MEMORY_CHARS = 15000  # keeps cached context (and its cost) bounded

app = Flask(
    __name__,
    static_folder=str(resource_path("static")),
    template_folder=str(resource_path("templates")),
)

# ---------- local settings (API keys entered via the Settings screen) ----------
# Stored in config.json inside DATA_DIR — the same permanent, per-user location as
# everything else. .env still works unchanged for development; config.json takes priority
# when both are present, since it reflects what the person actually entered in the app.

CONFIG_PATH = DATA_DIR / "config.json"


def load_config() -> dict:
    if CONFIG_PATH.exists():
        try:
            return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_config(updates: dict) -> dict:
    current = load_config()
    current.update(updates)
    CONFIG_PATH.write_text(json.dumps(current, indent=2), encoding="utf-8")
    return current


def get_api_key(config_key: str, env_var: str) -> str | None:
    return load_config().get(config_key) or os.environ.get(env_var)


def reset_clients():
    """Called after Settings saves a new key, so the next API call picks it up instead of
    reusing an already-created client built from the old (or missing) key."""
    global _client, _genai_client
    _client = None
    _genai_client = None


_client = None


def get_client():
    """Lazily create the Anthropic client so a missing key doesn't crash startup."""
    global _client
    if _client is None:
        from anthropic import Anthropic

        api_key = get_api_key("anthropic_api_key", "ANTHROPIC_API_KEY")
        if not api_key:
            raise RuntimeError(
                "No Anthropic API key set. Open Settings and paste your key, or set "
                "ANTHROPIC_API_KEY in .env for development."
            )
        _client = Anthropic(api_key=api_key)
    return _client


_genai_client = None


def get_genai_client():
    """Lazily create the Gemini client so a missing key doesn't crash startup."""
    global _genai_client
    if _genai_client is None:
        from google import genai

        api_key = get_api_key("google_api_key", "GOOGLE_API_KEY")
        if not api_key:
            raise RuntimeError(
                "No Google (Gemini) API key set. Open Settings and paste your key, or set "
                "GOOGLE_API_KEY in .env for development."
            )
        _genai_client = genai.Client(api_key=api_key)
    return _genai_client


# ---------- Excel topic queue ----------
# Expected header row (row 1): Day | Topic | Primary Keyword | Notes | Word Count | Status | Generated On
# Only Day, Topic, and Primary Keyword are needed — the rest are optional. "Status" and
# "Generated On" get written back into the same file after each generation, and the updated
# file can be downloaded from the queue card.

QUEUE_DIR = DATA_DIR / "data"
QUEUE_FILE = QUEUE_DIR / "seo_queue.xlsx"


def _header_map(ws):
    """Maps lowercased header name -> 1-based column index, from row 1."""
    mapping = {}
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value:
            mapping[str(cell.value).strip().lower()] = idx
    return mapping


def read_queue():
    """Reads the uploaded queue file and returns a list of row dicts."""
    if not QUEUE_FILE.exists():
        return []
    from openpyxl import load_workbook

    wb = load_workbook(QUEUE_FILE)
    ws = wb.active
    cols = _header_map(ws)

    def get(row_cells, name, default=""):
        idx = cols.get(name.lower())
        if idx is None:
            return default
        val = row_cells[idx - 1].value
        return default if val in (None, "") else str(val).strip()

    rows = []
    for r in range(2, ws.max_row + 1):
        row_cells = ws[r]
        topic = get(row_cells, "Topic")
        keyword = get(row_cells, "Primary Keyword")
        if not topic and not keyword:
            continue  # skip blank rows
        rows.append(
            {
                "row_number": r,
                "day": get(row_cells, "Day"),
                "topic": topic,
                "keyword": keyword,
                "notes": get(row_cells, "Notes"),
                "word_count": get(row_cells, "Word Count") or "800",
                "status": get(row_cells, "Status") or "Pending",
                "generated_on": get(row_cells, "Generated On"),
            }
        )
    return rows


def update_queue_status(row_number, status):
    """Writes Status + a timestamp back into the queue file for one row."""
    from openpyxl import load_workbook

    wb = load_workbook(QUEUE_FILE)
    ws = wb.active
    cols = _header_map(ws)
    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    if "status" in cols:
        ws.cell(row=row_number, column=cols["status"], value=status)
    if "generated on" in cols:
        ws.cell(row=row_number, column=cols["generated on"], value=timestamp)
    wb.save(QUEUE_FILE)
    return timestamp


def agent_dir(agent: str) -> Path:
    if agent not in ALLOWED_AGENTS:
        raise ValueError("unknown agent")
    d = MEMORY_DIR / agent
    d.mkdir(parents=True, exist_ok=True)
    return d


# ---------- draft persistence ----------
# Every generated draft is saved to disk as its own .md file (with a small YAML front-matter
# header) so nothing is lost on refresh, and so past drafts can be browsed and reopened.

DRAFTS_DIR = DATA_DIR / "drafts"


def slugify(text: str, max_len: int = 60) -> str:
    text = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return text[:max_len] or "untitled"


def save_draft(agent: str, topic: str, keyword: str, draft_text: str) -> dict:
    d = DRAFTS_DIR / agent
    d.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.datetime.now()
    filename = f"{timestamp.strftime('%Y%m%d-%H%M%S')}-{slugify(topic)}.md"

    title_match = re.search(r"^TITLE:\s*(.+)$", draft_text, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else topic
    generated_on = timestamp.strftime("%Y-%m-%d %H:%M")

    front_matter = (
        "---\n"
        f"topic: {topic}\n"
        f"keyword: {keyword}\n"
        f"title: {title}\n"
        f"generated_on: {generated_on}\n"
        "---\n\n"
    )
    (d / filename).write_text(front_matter + draft_text, encoding="utf-8")
    return {"filename": filename, "title": title, "generated_on": generated_on}


def _split_front_matter(raw: str):
    if not raw.startswith("---"):
        return {}, raw
    parts = raw.split("---", 2)
    if len(parts) < 3:
        return {}, raw
    fm_text, body = parts[1], parts[2].lstrip("\n")
    fm = {}
    for line in fm_text.strip().splitlines():
        if ":" in line:
            key, val = line.split(":", 1)
            fm[key.strip()] = val.strip()
    return fm, body


def read_draft(agent: str, filename: str) -> str:
    path = DRAFTS_DIR / agent / secure_filename(filename)
    if not path.exists():
        raise FileNotFoundError(filename)
    _, body = _split_front_matter(path.read_text(encoding="utf-8"))
    return body


def list_drafts(agent: str, limit: int = 50) -> list:
    d = DRAFTS_DIR / agent
    if not d.exists():
        return []
    files = sorted(d.glob("*.md"), key=lambda p: p.stat().st_mtime, reverse=True)[:limit]
    drafts = []
    for f in files:
        fm, _ = _split_front_matter(f.read_text(encoding="utf-8"))
        drafts.append(
            {
                "filename": f.name,
                "title": fm.get("title", f.stem),
                "generated_on": fm.get("generated_on", ""),
            }
        )
    return drafts


# ---------- Word archive ----------
# One running .docx per agent. Every generation gets appended automatically, newest entry
# first, separated by page breaks. Markdown headings/bold/bullets render as real Word
# formatting, not raw symbols.

ARCHIVE_DIR = DATA_DIR / "archive"


def archive_path(agent: str) -> Path:
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    return ARCHIVE_DIR / f"{agent}_archive.docx"


def _strip_title_meta_lines(draft_text: str) -> str:
    body = re.sub(r"^TITLE:.*$", "", draft_text, flags=re.MULTILINE)
    body = re.sub(r"^META DESCRIPTION:.*$", "", body, flags=re.MULTILINE)
    return body.strip()


def _new_paragraph_before(doc, anchor):
    """Creates a paragraph positioned immediately before `anchor` (or appended to the end
    of the document if anchor is None) and returns it."""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    if anchor is None:
        return doc.add_paragraph()
    p_element = OxmlElement("w:p")
    anchor._p.addprevious(p_element)
    return Paragraph(p_element, anchor._parent)


def _add_inline_runs(paragraph, text):
    """Adds text to a paragraph, turning **bold** markdown into an actual bold run."""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def append_to_archive(agent: str, topic: str, keyword: str, title: str, generated_on: str, body_markdown: str):
    """Adds a new entry to the top of the agent's Word archive, creating the file on first
    use. Newest entry always sits directly under the top-level title; older entries are
    pushed down and separated with a page break."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    path = archive_path(agent)
    doc = Document(str(path)) if path.exists() else Document()
    if not path.exists():
        doc.add_heading(f"{agent.upper()} content archive", level=0)

    paragraphs = doc.paragraphs
    anchor = paragraphs[1] if len(paragraphs) > 1 else None  # previous newest entry, if any

    def insert_p():
        return _new_paragraph_before(doc, anchor)

    title_p = insert_p()
    title_p.style = doc.styles["Heading 1"]
    title_p.add_run(title or topic)

    meta_p = insert_p()
    meta_run = meta_p.add_run(f"Topic: {topic}   |   Keyword: {keyword}   |   Generated: {generated_on}")
    meta_run.italic = True
    meta_run.font.size = Pt(9)

    for line in body_markdown.split("\n"):
        line = line.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            p = insert_p()
            p.style = doc.styles["Heading 3"]
            p.add_run(line[4:].strip())
        elif line.startswith("## "):
            p = insert_p()
            p.style = doc.styles["Heading 2"]
            p.add_run(line[3:].strip())
        elif line.strip().startswith(("- ", "* ")):
            p = insert_p()
            p.style = doc.styles["List Bullet"]
            _add_inline_runs(p, line.strip()[2:])
        else:
            p = insert_p()
            _add_inline_runs(p, line.strip())

    if anchor is not None:
        break_p = insert_p()
        break_p.add_run().add_break(WD_BREAK.PAGE)

    doc.save(str(path))


def extract_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    return path.read_text(encoding="utf-8", errors="ignore")


def build_memory_block(agent: str) -> tuple[str, bool]:
    """Concatenate every file in an agent's memory folder into one text block.
    Returns (text, truncated)."""
    d = agent_dir(agent)
    chunks = []
    for f in sorted(d.iterdir()):
        if f.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        try:
            text = extract_text(f).strip()
        except Exception as e:
            text = f"[Could not read {f.name}: {e}]"
        if text:
            chunks.append(f"### {f.name}\n{text}")
    combined = "\n\n".join(chunks)
    truncated = len(combined) > MAX_MEMORY_CHARS
    return combined[:MAX_MEMORY_CHARS], truncated


# ---------- Reddit/Quora safety checker ----------
# A second Claude call reviews every draft against the same rules in the memory folder.
# Runs automatically after generation; never auto-rewrites anything — just reports issues
# and logs them to a dedicated, human-reviewable file inside memory.

LESSONS_FILENAME = "_lessons_learned.md"

SAFETY_CHECK_INSTRUCTIONS = """You are an experienced Reddit and Quora moderator and community \
veteran, reviewing this draft before it gets posted somewhere on those platforms (or a guest \
site with similar norms). You've seen thousands of posts get removed or quietly down-ranked for \
self-promotion, and you know it's rarely just banned words — it's the overall feel of the piece.

Judge it the way a skeptical, experienced moderator would, across all of these:
- Self-promotional ratio: does this read as genuinely useful content that happens to mention a \
brand, or as an ad wearing a helpful-content costume? Zero banned words can still fail this.
- AI-generated tells: generic phrasing, over-structured listicle formatting, repetitive sentence \
rhythm, keyword-stuffed headers, or a title/intro that oversells — the kind of thing that gets a \
post reported before anyone even reads closely.
- Whether any specific claim would survive someone asking "source?" in the comments.
- Whether the piece still makes complete sense to a reader who never noticed the brand mention.
- The brand's own hard rules, provided separately below — these are the mandatory minimum, not \
the whole bar. Check them specifically, then judge everything else above using your own read.

Be genuinely strict. The real question is: if a wary moderator or a suspicious commenter read \
this closely, would it survive? If you're not confident it would, flag it — don't rubber-stamp \
anything you're on the fence about.

Respond in exactly this format, nothing else, no preamble:

If the draft would survive that scrutiny:
STATUS: PASS
NOTE: <one short, specific sentence about what it did well>

If there are issues, however small:
STATUS: NEEDS_FIXES
ISSUE 1: <what's wrong — quote the exact phrase, or describe the structural/tonal issue>
FIX 1: <the specific rewritten replacement text, or a specific structural fix>
ISSUE 2: <next issue, same format>
FIX 2: <...>
(list every issue found — don't stop at the first one)
"""


def run_safety_check(agent: str, draft_text: str) -> str:
    memory_text, _ = build_memory_block(agent)
    memory_block = memory_text if memory_text else "No brand guidelines uploaded yet."
    client = get_client()
    response = client.messages.create(
        model="claude-sonnet-5",
        max_tokens=1200,
        system=[
            {"type": "text", "text": SAFETY_CHECK_INSTRUCTIONS},
            {
                "type": "text",
                "text": f"Brand guidelines to check against:\n\n{memory_block}",
                "cache_control": {"type": "ephemeral"},
            },
        ],
        messages=[{"role": "user", "content": f"Review this draft:\n\n{draft_text}"}],
    )
    return "".join(block.text for block in response.content if block.type == "text")


def parse_safety_check(text: str) -> dict:
    text = text.strip()
    status_match = re.search(r"STATUS:\s*(PASS|NEEDS_FIXES)", text, re.IGNORECASE)
    status = status_match.group(1).upper() if status_match else "NEEDS_FIXES"

    if status == "PASS":
        note_match = re.search(r"NOTE:\s*(.+)", text)
        return {"status": "pass", "note": note_match.group(1).strip() if note_match else "", "issues": []}

    issues = []
    for issue, fix in re.findall(
        r"ISSUE\s*\d+:\s*(.+?)\nFIX\s*\d+:\s*(.+?)(?=\n?ISSUE\s*\d+:|\Z)", text, re.DOTALL
    ):
        issues.append({"issue": issue.strip(), "fix": fix.strip()})
    return {"status": "needs_fixes", "note": "", "issues": issues}


def append_lesson(agent: str, issue: str, fix: str):
    """Appends a flagged issue to a dedicated, auto-included memory file — visible and
    prunable, rather than silently rewriting the brand rules the checker judges against."""
    d = agent_dir(agent)
    path = d / LESSONS_FILENAME
    if not path.exists():
        path.write_text(
            "# Lessons learned (auto-logged by the safety checker)\n\n"
            "Flagged issues get added here automatically. Review periodically — prune what "
            "isn't useful, and promote recurring patterns into your main brand rules file.\n\n",
            encoding="utf-8",
        )
    date = datetime.datetime.now().strftime("%Y-%m-%d")
    with path.open("a", encoding="utf-8") as f:
        f.write(f"- [{date}] {issue.strip()} → {fix.strip()}\n")


# ---------- settings endpoints ----------

@app.route("/api/settings")
def get_settings():
    cfg = load_config()
    # Never send the actual key values back to the browser — only whether each is set.
    return jsonify(
        anthropic_api_key_set=bool(cfg.get("anthropic_api_key") or os.environ.get("ANTHROPIC_API_KEY")),
        google_api_key_set=bool(cfg.get("google_api_key") or os.environ.get("GOOGLE_API_KEY")),
    )


@app.route("/api/settings", methods=["POST"])
def update_settings():
    data = request.get_json(force=True) or {}
    updates = {}
    if data.get("anthropic_api_key"):
        updates["anthropic_api_key"] = data["anthropic_api_key"].strip()
    if data.get("google_api_key"):
        updates["google_api_key"] = data["google_api_key"].strip()
    if updates:
        save_config(updates)
        reset_clients()
    return jsonify(ok=True)


# ---------- static pages ----------

@app.route("/")
def index():
    return send_from_directory("templates", "index.html")


# ---------- memory endpoints ----------

@app.route("/api/memory/<agent>/list")
def memory_list(agent):
    try:
        d = agent_dir(agent)
    except ValueError:
        return jsonify(error="unknown agent"), 404
    files = []
    for f in sorted(d.iterdir()):
        if f.suffix.lower() in ALLOWED_EXTENSIONS:
            files.append({"name": f.name, "size": f.stat().st_size})
    return jsonify(files=files)


@app.route("/api/memory/<agent>/upload", methods=["POST"])
def memory_upload(agent):
    try:
        d = agent_dir(agent)
    except ValueError:
        return jsonify(error="unknown agent"), 404
    if "file" not in request.files:
        return jsonify(error="no file provided"), 400
    file = request.files["file"]
    filename = secure_filename(file.filename or "")
    ext = Path(filename).suffix.lower()
    if not filename or ext not in ALLOWED_EXTENSIONS:
        return jsonify(error="only .txt, .md, and .pdf files are supported"), 400
    file.save(d / filename)
    return jsonify(ok=True, name=filename)


@app.route("/api/memory/<agent>/<filename>", methods=["DELETE"])
def memory_delete(agent, filename):
    try:
        d = agent_dir(agent)
    except ValueError:
        return jsonify(error="unknown agent"), 404
    target = d / secure_filename(filename)
    if target.exists():
        target.unlink()
        return jsonify(ok=True)
    return jsonify(error="not found"), 404


# ---------- Excel queue endpoints ----------

@app.route("/api/queue/upload", methods=["POST"])
def queue_upload():
    if "file" not in request.files:
        return jsonify(error="no file provided"), 400
    file = request.files["file"]
    filename = secure_filename(file.filename or "")
    if not filename.lower().endswith(".xlsx"):
        return jsonify(error="please upload a .xlsx file"), 400

    QUEUE_DIR.mkdir(parents=True, exist_ok=True)
    file.save(QUEUE_FILE)

    try:
        rows = read_queue()
    except Exception as e:
        return jsonify(error=f"Could not read that file: {e}"), 400
    return jsonify(ok=True, rows=rows)


@app.route("/api/queue/rows")
def queue_rows():
    try:
        rows = read_queue()
    except Exception as e:
        return jsonify(error=f"Could not read the queue file: {e}"), 500
    return jsonify(rows=rows, has_file=QUEUE_FILE.exists())


@app.route("/api/queue/download")
def queue_download():
    from flask import send_file

    if not QUEUE_FILE.exists():
        return jsonify(error="no file uploaded yet"), 404
    return send_file(QUEUE_FILE, as_attachment=True, download_name="seo_queue.xlsx")


# ---------- draft endpoints ----------

@app.route("/api/drafts/<agent>")
def drafts_list(agent):
    if agent not in ALLOWED_AGENTS:
        return jsonify(error="unknown agent"), 404
    return jsonify(drafts=list_drafts(agent))


@app.route("/api/drafts/<agent>/<filename>")
def drafts_get(agent, filename):
    if agent not in ALLOWED_AGENTS:
        return jsonify(error="unknown agent"), 404
    try:
        body = read_draft(agent, filename)
    except FileNotFoundError:
        return jsonify(error="draft not found"), 404
    return jsonify(draft=body)


@app.route("/api/archive/<agent>/download")
def archive_download(agent):
    from flask import send_file

    if agent not in ALLOWED_AGENTS:
        return jsonify(error="unknown agent"), 404
    path = archive_path(agent)
    if not path.exists():
        return jsonify(error="archive is empty — generate a draft first"), 404
    return send_file(path, as_attachment=True, download_name=f"{agent}_archive.docx")


# ---------- image generation ----------
# Reference photos live in a folder you manage yourself in File Explorer — no upload
# mechanism at all. Set IMAGE_REF_ROOT in .env to a folder containing sierra/, malibu/, and
# laguna/ subfolders; if it's not set, it defaults to memory/image/ inside this project.
# Generated outputs are separate and still fully app-managed under memory/image/<product>/output/.

IMAGE_MODEL = "gemini-2.5-flash-image"
PRODUCTS = {"sierra", "malibu", "laguna"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
IMAGE_OUTPUT_ROOT = MEMORY_DIR / "image"
IMAGE_REF_ROOT = Path(os.environ["IMAGE_REF_ROOT"]) if os.environ.get("IMAGE_REF_ROOT") else IMAGE_OUTPUT_ROOT


def product_ref_dir(product: str) -> Path:
    if product not in PRODUCTS:
        raise ValueError("unknown product")
    d = IMAGE_REF_ROOT / product
    d.mkdir(parents=True, exist_ok=True)
    return d


def product_output_dir(product: str) -> Path:
    if product not in PRODUCTS:
        raise ValueError("unknown product")
    d = IMAGE_OUTPUT_ROOT / product / "output"
    d.mkdir(parents=True, exist_ok=True)
    return d


def load_reference_images(product: str) -> list:
    from PIL import Image

    d = product_ref_dir(product)
    images = []
    for f in sorted(d.iterdir()):
        if f.suffix.lower() in IMAGE_EXTENSIONS:
            images.append(Image.open(f))
    return images


def extract_image_and_text(response):
    """Pulls the generated image bytes and any accompanying text note out of a Gemini
    image-model response."""
    image_bytes = None
    text_note = ""
    for part in response.parts:
        if getattr(part, "text", None):
            text_note += part.text
        elif getattr(part, "inline_data", None) is not None:
            image_bytes = part.inline_data.data
    return image_bytes, text_note.strip()


def save_output_image(product: str, image_bytes: bytes) -> str:
    """Saves a generated image into that product's output folder with a unique,
    timestamp-based filename, and returns just the filename."""
    d = product_output_dir(product)
    filename = f"{datetime.datetime.now().strftime('%Y%m%d-%H%M%S-%f')}.png"
    (d / filename).write_bytes(image_bytes)
    return filename


def list_output_images(product: str, limit: int = 100) -> list:
    d = product_output_dir(product)
    files = sorted(d.glob("*.png"), key=lambda p: p.stat().st_mtime, reverse=True)[:limit]
    return [
        {
            "filename": f.name,
            "generated_on": datetime.datetime.fromtimestamp(f.stat().st_mtime).strftime(
                "%Y-%m-%d %H:%M"
            ),
        }
        for f in files
    ]


# ---------- reference photo endpoints (read-only — files live in IMAGE_REF_ROOT) ----------

@app.route("/api/image-refs/<product>/list")
def image_refs_list(product):
    try:
        d = product_ref_dir(product)
    except ValueError:
        return jsonify(error="unknown product"), 404
    files = [f.name for f in sorted(d.iterdir()) if f.suffix.lower() in IMAGE_EXTENSIONS]
    return jsonify(files=files, folder=str(d))


@app.route("/api/image-refs/<product>/file/<filename>")
def image_refs_file(product, filename):
    try:
        d = product_ref_dir(product)
    except ValueError:
        return jsonify(error="unknown product"), 404
    safe_name = secure_filename(filename)
    if not (d / safe_name).exists():
        return jsonify(error="not found"), 404
    return send_from_directory(d, safe_name)


# ---------- output gallery endpoints ----------

@app.route("/api/image-outputs/<product>/list")
def image_outputs_list(product):
    try:
        outputs = list_output_images(product)
    except ValueError:
        return jsonify(error="unknown product"), 404
    return jsonify(outputs=outputs)


@app.route("/api/image-outputs/<product>/file/<filename>")
def image_outputs_file(product, filename):
    try:
        d = product_output_dir(product)
    except ValueError:
        return jsonify(error="unknown product"), 404
    safe_name = secure_filename(filename)
    if not (d / safe_name).exists():
        return jsonify(error="not found"), 404
    return send_from_directory(d, safe_name)


# ---------- generation + conversational editing ----------

@app.route("/api/generate/image", methods=["POST"])
def generate_image():
    data = request.get_json(force=True) or {}
    product = (data.get("product") or "").strip().lower()
    prompt = (data.get("prompt") or "").strip()

    if product not in PRODUCTS:
        return jsonify(error="a valid product (sierra, malibu, or laguna) is required"), 400
    if not prompt:
        return jsonify(error="a description is required"), 400

    try:
        client = get_genai_client()
    except RuntimeError as e:
        return jsonify(error=str(e)), 500

    try:
        ref_images = load_reference_images(product)
    except Exception as e:
        return jsonify(error=f"Could not load reference photos: {e}"), 500

    full_prompt = (
        f"Generate a photorealistic image: {prompt}\n\n"
        f"Reference photos of White Pebble Wonders' {product.title()} pebbles are attached — "
        "match their real size, shape, and color as closely as possible rather than "
        "generic decorative stones."
    )
    contents = [full_prompt] + ref_images

    try:
        response = client.models.generate_content(model=IMAGE_MODEL, contents=contents)
    except Exception as e:
        return jsonify(error=f"Image generation error: {e}"), 502

    image_bytes, note = extract_image_and_text(response)
    if not image_bytes:
        return jsonify(error="No image came back — try rephrasing the description."), 502

    filename = save_output_image(product, image_bytes)

    return jsonify(
        product=product,
        filename=filename,
        image_url=f"/api/image-outputs/{product}/file/{filename}",
        note=note,
    )


@app.route("/api/generate/image/edit", methods=["POST"])
def edit_image():
    from PIL import Image

    data = request.get_json(force=True) or {}
    product = (data.get("product") or "").strip().lower()
    filename = (data.get("filename") or "").strip()
    comment = (data.get("comment") or "").strip()

    if product not in PRODUCTS:
        return jsonify(error="a valid product (sierra, malibu, or laguna) is required"), 400
    if not comment:
        return jsonify(error="A comment describing the change is required"), 400

    try:
        d = product_output_dir(product)
    except ValueError:
        return jsonify(error="unknown product"), 404

    current_path = d / secure_filename(filename)
    if not current_path.exists():
        return jsonify(error="That image no longer exists — generate a new one first."), 404

    try:
        client = get_genai_client()
    except RuntimeError as e:
        return jsonify(error=str(e)), 500

    current_image = Image.open(current_path)

    edit_prompt = (
        f"Edit this image: {comment}\n\n"
        "Keep everything else about the image the same — only change what was asked."
    )
    contents = [edit_prompt, current_image]

    try:
        response = client.models.generate_content(model=IMAGE_MODEL, contents=contents)
    except Exception as e:
        return jsonify(error=f"Image edit error: {e}"), 502

    image_bytes, note = extract_image_and_text(response)
    if not image_bytes:
        return jsonify(error="No image came back — try rephrasing the edit."), 502

    new_filename = save_output_image(product, image_bytes)

    return jsonify(
        product=product,
        filename=new_filename,
        image_url=f"/api/image-outputs/{product}/file/{new_filename}",
        note=note,
    )


# ---------- generation endpoints ----------

SEO_SYSTEM_INSTRUCTIONS = """You are a senior SEO content writer producing off-page content \
(guest posts, partner-site articles, and blog content) for a small manufacturing/D2C business.

Follow these rules on every draft:
- Put the primary keyword naturally in the title, the first paragraph, and at least one H2.
- Use short paragraphs (2-4 sentences), scannable H2/H3 subheadings, and at least one list where it helps readability.
- Write for a human reader first; never stuff keywords or pad sentences to hit a word count.
- Match the brand voice, terminology, and any hard rules described in the "Brand guidelines" \
section below, if one is provided. If none is provided, default to a clear, confident, plain-English tone.
- Output in this exact structure:

TITLE: <title>
META DESCRIPTION: <150-160 character meta description>

<body in markdown, using ## for H2 and ### for H3>
"""


@app.route("/api/generate/seo", methods=["POST"])
def generate_seo():
    data = request.get_json(force=True) or {}
    topic = (data.get("topic") or "").strip()
    keyword = (data.get("keyword") or "").strip()
    word_count = (data.get("word_count") or "800").strip()
    notes = (data.get("notes") or "").strip()
    row_number = data.get("row_number")  # present when triggered from the Excel queue

    if not topic or not keyword:
        return jsonify(error="topic and target keyword are required"), 400

    memory_text, truncated = build_memory_block("seo")
    memory_block = memory_text if memory_text else "No brand guidelines uploaded yet."

    user_brief = (
        f"Topic / working title: {topic}\n"
        f"Primary target keyword: {keyword}\n"
        f"Target length: ~{word_count} words\n"
        f"Additional notes: {notes or 'none'}"
    )

    try:
        client = get_client()
        response = client.messages.create(
            model="claude-sonnet-5",
            max_tokens=3000,
            system=[
                {"type": "text", "text": SEO_SYSTEM_INSTRUCTIONS},
                {
                    "type": "text",
                    "text": f"Brand guidelines (from uploaded memory files):\n\n{memory_block}",
                    "cache_control": {"type": "ephemeral"},
                },
            ],
            messages=[{"role": "user", "content": user_brief}],
        )
    except RuntimeError as e:
        return jsonify(error=str(e)), 500
    except Exception as e:
        return jsonify(error=f"Claude API error: {e}"), 502

    draft = "".join(block.text for block in response.content if block.type == "text")

    saved = save_draft("seo", topic, keyword, draft)

    try:
        append_to_archive(
            "seo", topic, keyword, saved["title"], saved["generated_on"],
            _strip_title_meta_lines(draft),
        )
    except Exception:
        # A failed archive write shouldn't cost the person their draft — the .md file
        # already saved it either way.
        pass

    safety = {"status": "unknown", "note": "", "issues": []}
    try:
        safety_raw = run_safety_check("seo", draft)
        safety = parse_safety_check(safety_raw)
        if safety["status"] == "needs_fixes":
            for item in safety["issues"]:
                try:
                    append_lesson("seo", item["issue"], item["fix"])
                except Exception:
                    pass
    except Exception:
        # A failed safety check shouldn't block the draft itself from coming back.
        pass

    queue_status_updated = False
    if row_number:
        try:
            update_queue_status(int(row_number), "Generated")
            queue_status_updated = True
        except Exception:
            # The draft is still good even if the file write fails — surface it via the flag
            # rather than failing the whole request.
            pass

    return jsonify(
        draft=draft,
        memory_truncated=truncated,
        queue_status_updated=queue_status_updated,
        saved_draft=saved,
        safety_check=safety,
        model="claude-sonnet-5",
        generated_at=int(time.time()),
        usage={
            "input_tokens": response.usage.input_tokens,
            "output_tokens": response.usage.output_tokens,
            "cache_read_input_tokens": getattr(response.usage, "cache_read_input_tokens", 0),
            "cache_creation_input_tokens": getattr(response.usage, "cache_creation_input_tokens", 0),
        },
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)